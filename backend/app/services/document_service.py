import io
from typing import List
from pypdf import PdfReader
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select

from backend.app.models import Document
from backend.app.services.qdrant_service import QdrantService

class DocumentService:
    def __init__(self, qdrant_service: QdrantService):
        self.qdrant_service = qdrant_service
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=700,
            chunk_overlap=150,
            length_function=len
        )

    def extract_text_from_pdf(self, file_bytes: bytes) -> str:
        pdf_file = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_file)
        text_parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n".join(text_parts)

    def extract_text_from_docx(self, file_bytes: bytes) -> str:
        docx_file = io.BytesIO(file_bytes)
        doc = DocxDocument(docx_file)
        text_parts = []
        for para in doc.paragraphs:
            if para.text:
                text_parts.append(para.text)
        return "\n".join(text_parts)

    async def ingest_document(
        self,
        db: AsyncSession,
        filename: str,
        category: str,
        file_bytes: bytes,
        user_id: int
    ) -> Document:
        # Create Document record in PostgreSQL with status 'processing'
        db_doc = Document(
            filename=filename,
            category=category.lower(),
            uploaded_by=user_id,
            status="processing"
        )
        db.add(db_doc)
        await db.commit()
        await db.refresh(db_doc)

        try:
            # Extract text depending on file extension
            ext = filename.split(".")[-1].lower()
            if ext == "pdf":
                text = self.extract_text_from_pdf(file_bytes)
            elif ext in ["docx", "doc"]:
                text = self.extract_text_from_docx(file_bytes)
            elif ext == "txt":
                text = file_bytes.decode("utf-8", errors="ignore")
            else:
                raise ValueError(f"Unsupported file extension: {ext}")

            if not text.strip():
                raise ValueError("Extracted text is empty")

            # Chunk the text
            chunks = self.text_splitter.split_text(text)

            # Upload vectors to Qdrant
            qdrant_success = await self.qdrant_service.add_document_chunks(
                document_id=db_doc.id,
                filename=db_doc.filename,
                category=db_doc.category,
                chunks=chunks
            )

            if qdrant_success:
                db_doc.status = "completed"
            else:
                db_doc.status = "failed"

        except Exception as e:
            print(f"Error ingesting document {filename}: {e}")
            db_doc.status = "failed"

        await db.commit()
        await db.refresh(db_doc)
        return db_doc

    async def delete_document(self, db: AsyncSession, document_id: int) -> bool:
        # Fetch document
        result = await db.execute(select(Document).where(Document.id == document_id))
        db_doc = result.scalars().first()
        if not db_doc:
            return False

        try:
            # Delete vectors from Qdrant
            await self.qdrant_service.delete_document_vectors(document_id)
            
            # Delete metadata record from Postgres
            await db.delete(db_doc)
            await db.commit()
            return True
        except Exception as e:
            print(f"Error deleting document {document_id}: {e}")
            return False
